"""
The history as the BFP fire log lays it out (GET /api/history/export/), as
an Excel file or a CSV: the log's three header rows, with the same merged
cells, words and colours, then one record per row in columns A-W.

  A-B   date, and time of alarm as the number HHMM shown as HH:MM
  C-E   Region_7, Cebu, Talisay_City_c
  F     the exact location: the address as typed, or the record's puroks
        and barangay (bfp.exact_location)
  G-H   a fire's property type; for any other hazard, its type and sub
        category, as the flood and landslide log has them in G-H
  I-T   injured, then fatalities; firefighter, then civilian; male,
        female, total. The flood and landslide log counts people without
        either split, so its injured go in N and its dead in T.
  U-W   damage, affected structures, alarm level

A blank count stays blank: it means nobody counted, not zero.

The CSV is what Google Sheets saves from the log: the header rows with
merged cells left blank, dates as 26-Jan-23, times as 20:36, money as
Php75,000.00.

Plain functions over records, so they test without the database.
"""
import csv
import io
from datetime import time

from django.utils import timezone

from apps.history import bfp

BLUE, GREEN, CYAN = 'FF4FD1FF', 'FF92D050', 'FF3FCDFF'

# (cells, words, fill). Every cell of A1:W3 is in exactly one entry.
HEADER = [
    ('A1:A3', 'Date of Fire\n (date-month-year)', BLUE),
    ('B1:B3', 'Time of Alarm', GREEN),
    ('C1:C3', 'Region', GREEN),
    ('D1:D3', 'Province / District', GREEN),
    ('E1:E3', 'City / Municipality', GREEN),
    ('F1:F3', 'Exact Location', CYAN),
    ('G1:H2', 'Property Type', GREEN),
    ('G3', 'General Category', GREEN),
    ('H3', 'Sub Category', GREEN),
    ('I1:M1', 'Nr of Injured', GREEN),
    ('I2:J2', 'Firefighter', BLUE),
    ('I3', 'Male', BLUE),
    ('J3', 'Female', BLUE),
    ('K2:K3', 'Total', GREEN),
    ('L2:M2', 'Civilian', BLUE),
    ('L3', 'Male', BLUE),
    ('M3', 'Female', BLUE),
    # The log has the civilian injured total under Nr of Fatalities; kept so
    # the file matches it cell for cell
    ('N1:T1', 'Nr of Fatalities', GREEN),
    ('N2:N3', 'Total', GREEN),
    ('O2:P2', 'Firefighter', BLUE),
    ('O3', 'Male', BLUE),
    ('P3', 'Female', BLUE),
    ('Q2:Q3', 'Total', GREEN),
    ('R2:S2', 'Civilian', BLUE),
    ('R3', 'Male', BLUE),
    ('S3', 'Female', BLUE),
    ('T2:T3', 'Total', GREEN),
    ('U1:U3', 'Estimated Cost of Damage', BLUE),
    ('V1:V3', 'Nr of Affected Structure', BLUE),
    ('W1:W3', 'Alarm', GREEN),
]
HEADER_ROWS = 3
COLUMNS = 23  # A-W
WIDTHS = {'A': 18.6, 'B': 38.3, 'D': 20.6, 'E': 46.4, 'F': 34.7, 'G': 32.9, 'H': 28.9, 'I': 8.7,
          'U': 23.3, 'V': 8.7}
FORMATS = {'A': 'd-mmm-yy', 'B': '00\\:00', 'F': '@', 'G': '@', 'U': '"Php"#,##0.00',
           **{c: '#,##0' for c in 'IJKLMNOPQRSTV'}}
TOTALS = 'KNQT'  # bold, as in the log
WRAPPED = 'DEFHIJKLMNOPQRSTW'
REGION, PROVINCE, CITY = 'Region_7', 'Cebu', 'Talisay_City_c'


def _total(*values):
    known = [v for v in values if v is not None]
    return sum(known) if known else None


def _is_fire(record):
    return bool(record.hazard_type) and record.hazard_type.name.lower() == 'fire'


def _counts(record, fire):
    """Columns I-T."""
    if fire:
        ffm, fff, cm, cf, dffm, dfff, dcm, dcf = (getattr(record, k) for k in bfp.CASUALTY_FIELDS)
        injured, dead = _total(cm, cf), _total(dcm, dcf)
        known = _total(ffm, fff, cm, cf, dffm, dfff, dcm, dcf)
    else:
        ffm = fff = cm = cf = dffm = dfff = dcm = dcf = None
        injured, dead = record.injured, record.fatalities
        known = _total(injured, dead)
    # Casualties the record has no split for go in T, where the log puts
    # them (bfp.total_casualties reads T as everyone hurt or killed)
    unsplit = (record.total_casualties or 0) - (known or 0)
    if unsplit > 0:
        dead = (dead or 0) + unsplit
    return [ffm, fff, _total(ffm, fff), cm, cf, injured, dffm, dfff, _total(dffm, dfff), dcm, dcf, dead]


def log_row(record):
    """One record as columns A-W: dates, numbers and text, None where blank."""
    when = timezone.localtime(record.occurred_at, bfp.MANILA)
    fire = _is_fire(record)
    location = record.location_text or bfp.exact_location(
        [p.name for p in record.puroks.all()], record.barangay.name if record.barangay else None)
    if fire:
        kind, sub = record.property_category, record.property_subcategory
    else:
        kind, sub = (record.hazard_type.name.upper() if record.hazard_type else ''), record.hazard_subcategory
    return [
        when.date(), when.hour * 100 + when.minute, REGION, PROVINCE, CITY, location, kind, sub,
        *_counts(record, fire),
        record.damage_cost, record.structures_affected, record.alarm_level,
    ]


def to_xlsx(rows):
    """The .xlsx file's bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter, range_boundaries

    wb = Workbook()
    ws = wb.active
    ws.title = 'History'
    thin = Side(style='thin')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cells, words, fill in HEADER:
        first_col, first_row, last_col, last_row = range_boundaries(cells)
        for r in range(first_row, last_row + 1):
            for c in range(first_col, last_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = Font(name='Arial', size=8, bold=True)
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.fill = PatternFill('solid', fgColor=fill)
                cell.border = border
        ws.cell(row=first_row, column=first_col, value=words)
        if ':' in cells:
            ws.merge_cells(cells)

    for r, values in enumerate(rows, start=HEADER_ROWS + 1):
        for c, value in enumerate(values, start=1):
            letter = get_column_letter(c)
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str) and value.startswith('='):
                cell.data_type = 's'  # an address, never a formula
            cell.font = Font(name='Arial', size=10, bold=letter in TOTALS)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=letter in WRAPPED)
            cell.border = border
            if letter in FORMATS:
                cell.number_format = FORMATS[letter]

    for letter, width in WIDTHS.items():
        ws.column_dimensions[letter].width = width

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _csv_cell(value, letter):
    if value is None:
        return ''
    if letter == 'A':
        return f'{value.day}-{value:%b-%y}'
    if letter == 'B':
        return time(value // 100, value % 100).strftime('%H:%M')
    if letter == 'U':
        return f'Php{value:,.2f}'
    text = str(value)
    # A spreadsheet reads a leading = + - @ as a formula (as utils/csv.js does)
    return f"'{text}" if isinstance(value, str) and text[:1] in ('=', '+', '-', '@', '\t', '\r') else text


def to_csv(rows):
    """The CSV text, with a BOM so Excel reads it as UTF-8 (ñ in place names)."""
    from openpyxl.utils import get_column_letter, range_boundaries

    header = [[''] * COLUMNS for _ in range(HEADER_ROWS)]
    for cells, words, _ in HEADER:
        col, row, _, _ = range_boundaries(cells)
        header[row - 1][col - 1] = words
    letters = [get_column_letter(c) for c in range(1, COLUMNS + 1)]

    out = io.StringIO()
    writer = csv.writer(out, lineterminator='\r\n')
    writer.writerows(header)
    writer.writerows([_csv_cell(v, letter) for v, letter in zip(values, letters)] for values in rows)
    return '﻿' + out.getvalue()
