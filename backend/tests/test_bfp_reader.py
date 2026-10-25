"""
Reading the BFP fire log: one row per fire, below three header rows.

The sheet is typed by hand, so the dates come as Excel serials or as text
("JULY 16,2023" with the time "1158H/16 JULY 2023"), and a blank cell means
nobody filled it in, which is not the same as zero.
"""
import tempfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

from django.test import SimpleTestCase
from openpyxl import Workbook

from apps.history import bfp

MNL = ZoneInfo('Asia/Manila')


class SheetTestCase(SimpleTestCase):
    def setUp(self):
        self._dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._dir.cleanup)

    def make_sheet(self, rows):
        """An .xlsx with three header rows, then each {column letter: value} from row 4."""
        wb = Workbook()
        ws = wb.active
        ws['A1'], ws['I1'], ws['I2'], ws['I3'] = 'Date of Fire', 'Nr of Injured', 'Firefighter', 'Male'
        for i, cells in enumerate(rows, start=4):
            for col, value in cells.items():
                ws[f'{col}{i}'] = value
        path = Path(self._dir.name) / 'fires.xlsx'
        wb.save(path)
        return path


class ReadSheetTests(SheetTestCase):
    def test_reads_rows_below_the_three_header_rows(self):
        path = self.make_sheet([{'A': 44952, 'B': 2036, 'E': 'Talisay_City_c', 'F': 'PUROK SAMBAG, BRGY CAMP 4',
                                 'G': 'RESIDENTIAL', 'H': 'SINGLE AND TWO FAMILY DWELLING', 'I': 0, 'T': 0,
                                 'U': 75000, 'V': 1, 'W': '1st alarm '}])
        [row] = bfp.read_sheet(path)
        self.assertEqual(row['row'], 4)
        self.assertEqual(row['location'], 'PUROK SAMBAG, BRGY CAMP 4')
        self.assertEqual(row['damage_cost'], Decimal('75000'))
        self.assertEqual(row['structures_affected'], 1)
        self.assertEqual(row['alarm'], '1ST ALARM')
        self.assertEqual(row['injured_ff_male'], 0)
        self.assertIsNone(row['fatal_civ_female'])  # blank cell: not known

    def test_rows_without_a_location_are_skipped(self):
        self.assertEqual(bfp.read_sheet(self.make_sheet([{'A': 44952, 'B': 2036}])), [])

    def test_a_row_that_stops_early_reads_the_rest_as_blank(self):
        [row] = bfp.read_sheet(self.make_sheet([{'A': 44952, 'B': 2036, 'F': 'BRGY DUMLOG'}]))
        self.assertEqual(row['alarm'], '')
        self.assertIsNone(row['damage_cost'])
        self.assertIsNone(row['total_hurt'])


class ParseWhenTests(SimpleTestCase):
    def test_serial_dates_and_hhmm_times(self):
        self.assertEqual(bfp.parse_when(44952, 2036), datetime(2023, 1, 26, 20, 36, tzinfo=MNL))
        self.assertEqual(bfp.parse_when('44972.0', '952.0'), datetime(2023, 2, 15, 9, 52, tzinfo=MNL))

    def test_dates_the_sheet_already_formats_as_dates(self):
        # openpyxl hands back a datetime when the cell is formatted as a date
        self.assertEqual(bfp.parse_when(datetime(2026, 4, 29), 1045.0), datetime(2026, 4, 29, 10, 45, tzinfo=MNL))

    def test_text_dates_and_times(self):
        self.assertEqual(bfp.parse_when('JULY 16,2023', '1158H/16 JULY 2023'),
                         datetime(2023, 7, 16, 11, 58, tzinfo=MNL))
        self.assertEqual(bfp.parse_when('JULY 29,2023', '1130H/ 29 JULY 2023'),
                         datetime(2023, 7, 29, 11, 30, tzinfo=MNL))

    def test_blank_time_is_midnight_and_bad_date_is_none(self):
        self.assertEqual(bfp.parse_when(44952, None), datetime(2023, 1, 26, 0, 0, tzinfo=MNL))
        self.assertIsNone(bfp.parse_when('sometime', 1200))


class AlarmTests(SimpleTestCase):
    def test_alarm_levels(self):
        for alarm, code in [('FIRE OUT UPON ARRIVAL', 'Green'), ('UNRESPONDED', 'Green'), ('1ST ALARM', 'Green'),
                            ('2ND ALARM', 'Orange'), ('3RD ALARM', 'Red'), ('4TH ALARM', 'Red'),
                            ('TF ALPHA', 'Red'), ('GENERAL ALARM', 'Red'), (' 2nd alarm ', 'Orange')]:
            self.assertEqual(bfp.severity_for_alarm(alarm), (code, True), alarm)
        self.assertEqual(bfp.severity_for_alarm('ALARM X'), ('Green', False))


class CasualtyTests(SimpleTestCase):
    def test_casualties_take_the_larger_of_column_t_and_the_cells(self):
        zeros = {k: 0 for k in bfp.CASUALTY_FIELDS}
        self.assertEqual(bfp.total_casualties({**zeros, 'total_hurt': 4}), 4)
        self.assertEqual(bfp.total_casualties({**zeros, 'injured_civ_male': 1, 'total_hurt': 1}), 1)
        self.assertEqual(bfp.total_casualties({**zeros, 'injured_civ_male': 2, 'fatal_civ_male': 1, 'total_hurt': 1}), 3)
        self.assertIsNone(bfp.total_casualties({**{k: None for k in bfp.CASUALTY_FIELDS}, 'total_hurt': None}))


class RefTests(SimpleTestCase):
    def test_source_ref_ignores_spacing_and_case(self):
        when = datetime(2026, 4, 29, 10, 45, tzinfo=MNL)
        a = bfp.source_ref(when, 'Purok Sambag,  San Roque')
        self.assertEqual(a, bfp.source_ref(when, 'PUROK SAMBAG, SAN ROQUE'))
        self.assertTrue(a.startswith('bfp:'))
        self.assertEqual(len(a), 20)

    def test_only_talisay_rows_count(self):
        self.assertTrue(bfp.is_talisay({'city': 'Talisay_City_c'}))
        self.assertFalse(bfp.is_talisay({'city': 'Toledo_City'}))
