"""
GET /api/history/export/: the history in the BFP fire log's columns A-W, as
Excel or CSV, newest first. The rows and the files are checked without the
database; the view (filters, staff only, samples left out, the audit log)
against a real one.
"""
import csv
import io
from datetime import datetime, timedelta
from decimal import Decimal
from types import SimpleNamespace

from django.test import SimpleTestCase, TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.admin_dashboard.models import AuditLog
from apps.history import bfp
from apps.history.export import log_row, to_csv, to_xlsx
from apps.history.models import HistoricalRecord
from apps.history.views import HistoryExportView
from tests.fire_setup import FireSetupMixin

COUNTS = {k: None for k in bfp.CASUALTY_FIELDS}


def record(hazard='Fire', puroks=(), barangay='San Roque', **fields):
    """A HistoricalRecord stand-in: blank everything unless given."""
    values = {
        'occurred_at': datetime(2023, 2, 6, 22, 42, tzinfo=bfp.MANILA), 'location_text': '',
        'property_category': '', 'property_subcategory': '', 'hazard_subcategory': '',
        'injured': None, 'fatalities': None, 'total_casualties': None,
        'damage_cost': None, 'structures_affected': None, 'alarm_level': '', **COUNTS,
    }
    values.update(fields)
    return SimpleNamespace(
        hazard_type=SimpleNamespace(name=hazard) if hazard else None,
        barangay=SimpleNamespace(name=barangay) if barangay else None,
        puroks=SimpleNamespace(all=lambda: [SimpleNamespace(name=p) for p in puroks]),
        **values,
    )


def col(row, letter):
    return row[ord(letter) - ord('A')]


class LogRowTests(SimpleTestCase):
    def test_a_fire_fills_a_to_w_as_the_bfp_log_does(self):
        row = log_row(record(
            location_text='PUROK SAMBAG, LARAY, BRGY SAN ROQUE', property_category='RESIDENTIAL',
            property_subcategory='SINGLE AND TWO FAMILY DWELLING', damage_cost=Decimal('625000'),
            structures_affected=76, alarm_level='3RD ALARM', **{**COUNTS, 'injured_civ_male': 2, 'fatal_civ_female': 1},
        ))
        self.assertEqual(len(row), 23)
        self.assertEqual(row[:8], [datetime(2023, 2, 6).date(), 2242, 'Region_7', 'Cebu', 'Talisay_City_c',
                                   'PUROK SAMBAG, LARAY, BRGY SAN ROQUE', 'RESIDENTIAL',
                                   'SINGLE AND TWO FAMILY DWELLING'])
        self.assertEqual(row[8:20], [None, None, None, 2, None, 2, None, None, None, None, 1, 1])
        self.assertEqual(row[20:], [Decimal('625000'), 76, '3RD ALARM'])

    def test_the_date_and_time_are_manilas(self):
        row = log_row(record(occurred_at=datetime(2026, 9, 20, 10, 51, tzinfo=timezone.utc)))
        self.assertEqual((col(row, 'A'), col(row, 'B')), (datetime(2026, 9, 20).date(), 1851))

    def test_casualties_with_no_split_go_in_t(self):
        # The BFP log's 4 dead with no male/female cells filled, as imported
        row = log_row(record(total_casualties=4, **{k: 0 for k in bfp.CASUALTY_FIELDS}))
        self.assertEqual(row[8:20], [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 4])

    def test_a_flood_has_its_type_and_sub_category_in_g_h_and_its_counts_in_n_and_t(self):
        row = log_row(record('Flood', location_text='BRGY LINAO', hazard_subcategory='FLASH FLOOD',
                             injured=3, fatalities=1, total_casualties=4))
        self.assertEqual((col(row, 'G'), col(row, 'H')), ('FLOOD', 'FLASH FLOOD'))
        self.assertEqual(row[8:20], [None, None, None, None, None, 3, None, None, None, None, None, 1])
        self.assertEqual(row[20:], [None, None, ''])

    def test_with_no_address_the_puroks_and_barangay_are_the_location(self):
        row = log_row(record('Landslide', puroks=['Purok Banana'], barangay='Jaclupan'))
        self.assertEqual(col(row, 'F'), 'PUROK BANANA, BRGY JACLUPAN')


class FileTests(SimpleTestCase):
    def rows(self):
        return [log_row(record(location_text='=HYPERLINK("x")', property_category='RESIDENTIAL',
                               damage_cost=Decimal('75000'), **{**COUNTS, 'injured_ff_male': 1}))]

    def test_the_excel_file_has_the_logs_header_and_formats(self):
        ws = load_workbook(io.BytesIO(to_xlsx(self.rows()))).active
        self.assertIn('I1:M1', {str(m) for m in ws.merged_cells.ranges})
        self.assertEqual(len(ws.merged_cells.ranges), 20)
        self.assertEqual((ws['A1'].value, ws['G3'].value, ws['T2'].value, ws['W1'].value),
                         ('Date of Fire\n (date-month-year)', 'General Category', 'Total', 'Alarm'))
        self.assertEqual((ws['B4'].value, ws['B4'].number_format), (2242, '00\\:00'))
        self.assertEqual(ws['U4'].number_format, '"Php"#,##0.00')
        self.assertEqual((ws['K4'].value, ws['K4'].font.b, ws['I4'].font.b), (1, True, False))
        # An address is text, never a formula
        self.assertEqual((ws['F4'].value, ws['F4'].data_type), ('=HYPERLINK("x")', 's'))

    def test_the_csv_is_what_sheets_saves_from_the_log(self):
        text = to_csv(self.rows())
        self.assertTrue(text.startswith('﻿'))
        lines = list(csv.reader(io.StringIO(text[1:])))
        self.assertEqual(lines[0][:2], ['Date of Fire\n (date-month-year)', 'Time of Alarm'])
        self.assertEqual(lines[1][8:11], ['Firefighter', '', 'Total'])
        self.assertEqual(lines[2][6:8], ['General Category', 'Sub Category'])
        self.assertEqual(lines[3][:2], ['6-Feb-23', '22:42'])
        self.assertEqual((lines[3][5], lines[3][8], lines[3][9], lines[3][20]),
                         ("'=HYPERLINK(\"x\")", '1', '', 'Php75,000.00'))


class HistoryExportViewTests(FireSetupMixin, TestCase):
    def setUp(self):
        super().setUp()
        now = timezone.now()

        def make(hazard, days, **kw):
            return HistoricalRecord.objects.create(barangay=self.home, hazard_type=hazard, severity_level='Green',
                                                   occurred_at=now - timedelta(days=days), **kw)
        self.old_fire = make(self.fire, 30, location_text='OLD FIRE')
        self.new_fire = make(self.fire, 1, location_text='NEW FIRE')
        self.flood_rec = make(self.flood, 10, location_text='THE FLOOD')
        make(self.fire, 2, location_text='SAMPLE FIRE', is_sample=True)

    def get(self, user=None, **params):
        request = APIRequestFactory().get('/api/history/export/', params)
        if user:
            force_authenticate(request, user=user)
        return HistoryExportView.as_view()(request)

    def locations(self, response):
        ws = load_workbook(io.BytesIO(response.content)).active
        return [r[5] for r in ws.iter_rows(min_row=4, values_only=True)]

    def test_newest_first_with_no_sample_records(self):
        response = self.get(self.officer)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.locations(response), ['NEW FIRE', 'THE FLOOD', 'OLD FIRE'])

    def test_only_the_hazard_types_asked_for(self):
        response = self.get(self.officer, types='fire')
        self.assertEqual(self.locations(response), ['NEW FIRE', 'OLD FIRE'])
        self.assertIn('history-fire-', response['Content-Disposition'])

    def test_the_date_filter_applies(self):
        since = (timezone.now() - timedelta(days=15)).isoformat()
        self.assertEqual(self.locations(self.get(self.officer, date_from=since)), ['NEW FIRE', 'THE FLOOD'])

    def test_csv(self):
        response = self.get(self.officer, file='csv', types='Flood')
        self.assertEqual(response['Content-Type'], 'text/csv; charset=utf-8')
        self.assertRegex(response['Content-Disposition'], r'filename="history-flood-[\d-]+\.csv"')
        lines = list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))
        self.assertEqual([line[5] for line in lines[3:]], ['THE FLOOD'])

    def test_a_bad_file_kind_is_refused(self):
        self.assertEqual(self.get(self.officer, file='pdf').status_code, 400)

    def test_staff_only(self):
        self.assertIn(self.get().status_code, (401, 403))
        self.assertEqual(self.get(self.reporter).status_code, 403)

    def test_the_download_is_logged(self):
        self.get(self.officer, types='Fire')
        log = AuditLog.objects.get(action='EXPORT_REPORT')
        self.assertEqual(log.user, self.officer)
        self.assertIn('Excel: 2 record(s)', log.details)
        self.assertIn('types=Fire', log.details)
